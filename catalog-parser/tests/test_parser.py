import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from export_excel import save_excel, summary_rows
from parser import parse_cbr_rate, parse_listing, parse_product

FX = ROOT / "tests" / "fixtures"
BASE = "https://books.toscrape.com/catalogue/"


def read(name):
    return (FX / name).read_text(encoding="utf-8")


def test_listing_first_page():
    links, nxt = parse_listing(read("listing.html"), BASE + "page-1.html")
    assert len(links) == 20
    assert links[0] == BASE + "a-light-in-the-attic_1000/index.html"
    assert nxt == BASE + "page-2.html"


def test_listing_last_page_has_no_next():
    links, nxt = parse_listing(read("listing_last.html"), BASE + "page-50.html")
    assert len(links) == 20
    assert nxt is None


def test_product():
    url = BASE + "a-light-in-the-attic_1000/index.html"
    p = parse_product(read("product.html"), url, gbp_rate=100.0)
    assert p.title == "A Light in the Attic"
    assert p.category == "Poetry"
    assert p.price_gbp == 51.77
    assert p.price_rub == 5177.0
    assert p.in_stock == 22
    assert p.rating == 3
    assert p.upc == "a897fe39b1053632"
    assert p.description.startswith("It's hard to imagine")
    assert p.image.startswith("https://books.toscrape.com/media/")


def test_cbr_rate():
    rate = parse_cbr_rate((FX / "cbr.xml").read_bytes(), "GBP")
    assert 50 < rate < 300
    # номинал 100 (японская иена) делится корректно
    assert parse_cbr_rate((FX / "cbr.xml").read_bytes(), "JPY") < 5


def test_summary_and_excel(tmp_path):
    items = [
        dict(title="A", category="Poetry", price_rub=100.0, price_gbp=1.0, in_stock=2, rating=3,
             reviews=0, upc="1", description="d", image="http://x/1.jpg", url="http://x/1"),
        dict(title="B", category="Poetry", price_rub=300.0, price_gbp=3.0, in_stock=1, rating=5,
             reviews=0, upc="2", description="d", image="http://x/2.jpg", url="http://x/2"),
        dict(title="C", category="Travel", price_rub=50.0, price_gbp=0.5, in_stock=4, rating=1,
             reviews=0, upc="3", description="d", image="http://x/3.jpg", url="http://x/3"),
    ]
    assert summary_rows(items) == [["Poetry", 2, 200.0, 100.0, 300.0, 3], ["Travel", 1, 50.0, 50.0, 50.0, 4]]

    path = tmp_path / "out.xlsx"
    save_excel(items, 110.0, str(path))
    wb = load_workbook(path)
    assert wb.sheetnames == ["Товары", "Сводка"]
    ws = wb["Товары"]
    assert ws["A2"].value == "A" and ws["C2"].value == 100.0
    assert ws["H2"].hyperlink.target == "http://x/1"
    assert wb["Сводка"]["A4"].value == "Poetry"
