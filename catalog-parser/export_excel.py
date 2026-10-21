"""Выгрузка в Excel: лист с товарами + сводка по категориям с диаграммой."""
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [  # (ключ, заголовок, ширина, формат)
    ("title", "Название", 46, None),
    ("category", "Категория", 18, None),
    ("price_rub", "Цена, ₽", 12, '#,##0.00 "₽"'),
    ("price_gbp", "Цена, £", 10, '"£"0.00'),
    ("in_stock", "В наличии, шт", 13, "0"),
    ("rating", "Рейтинг", 10, "0"),
    ("upc", "Артикул (UPC)", 19, None),
    ("url", "Ссылка", 12, None),
    ("image", "Фото", 10, None),
    ("description", "Описание", 60, None),
]

HEAD_FILL = PatternFill("solid", fgColor="1D1B18")
HEAD_FONT = Font(bold=True, color="FFFFFF")


def _header(ws, titles):
    ws.append(titles)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def summary_rows(products: list[dict]) -> list[list]:
    """Сводка по категориям: кол-во, средняя/мин/макс цена, остаток на складе."""
    by_cat = defaultdict(list)
    for p in products:
        by_cat[p["category"]].append(p)
    rows = []
    for cat, items in by_cat.items():
        prices = [i["price_rub"] for i in items]
        rows.append([cat, len(items), round(sum(prices) / len(prices), 2), min(prices), max(prices),
                     sum(i["in_stock"] for i in items)])
    return sorted(rows, key=lambda r: -r[1])


def save_excel(products: list[dict], rate: float, path: str):
    wb = Workbook()

    # --- товары ---
    ws = wb.active
    ws.title = "Товары"
    _header(ws, [c[1] for c in COLUMNS])
    for p in products:
        ws.append([p[c[0]] for c in COLUMNS])
    n = len(products) + 1
    for i, (key, _, width, fmt) in enumerate(COLUMNS, 1):
        col = get_column_letter(i)
        ws.column_dimensions[col].width = width
        for cell in ws[col][1:]:
            if fmt:
                cell.number_format = fmt
            if key in ("url", "image") and cell.value:
                cell.hyperlink, cell.value, cell.style = cell.value, "открыть", "Hyperlink"
            if key == "description":
                cell.alignment = Alignment(wrap_text=False)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{n}"
    # цветовая шкала: дешёвое — зелёное, дорогое — красное
    ws.conditional_formatting.add(f"C2:C{n}", ColorScaleRule(
        start_type="min", start_color="C6EFCE", mid_type="percentile", mid_value=50,
        mid_color="FFF2CC", end_type="max", end_color="F8CBAD"))

    # --- сводка ---
    s = wb.create_sheet("Сводка")
    s.append([f"Выгрузка {datetime.now():%d.%m.%Y %H:%M} · товаров: {len(products)} · "
              f"курс ЦБ: 1 £ = {rate:.2f} ₽"])
    s["A1"].font = Font(italic=True, color="6D675E")
    s.append([])
    s.append(["Категория", "Товаров", "Средняя цена, ₽", "Мин. цена, ₽", "Макс. цена, ₽", "На складе, шт"])
    for c in s[3]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    rows = summary_rows(products)
    for r in rows:
        s.append(r)
    for col, w in zip("ABCDEF", (24, 10, 17, 14, 14, 14)):
        s.column_dimensions[col].width = w
    for row in s.iter_rows(min_row=4, min_col=3, max_col=5):
        for c in row:
            c.number_format = '#,##0 "₽"'

    if rows:
        top = min(len(rows), 15)
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Товаров по категориям (топ-15)"
        chart.legend = None
        chart.add_data(Reference(s, min_col=2, min_row=3, max_row=3 + top), titles_from_data=True)
        chart.set_categories(Reference(s, min_col=1, min_row=4, max_row=3 + top))
        chart.y_axis.majorGridlines = None
        chart.height, chart.width = 12, 18
        s.add_chart(chart, "H3")

    wb.save(path)
